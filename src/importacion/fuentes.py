"""Extracción de registros tabulares desde múltiples fuentes de datos."""

from __future__ import annotations

import csv
import json
from pathlib import Path
from typing import Any, Iterable, Mapping, Sequence


FuenteDatos = (
    str
    | Path
    | Mapping[str, Any]
    | Sequence[Mapping[str, Any]]
    | Iterable[Mapping[str, Any]]
    | Any
)


def extraer_registros(fuente: FuenteDatos, *, delimitador: str | None = None) -> list[dict[str, Any]]:
    """Normaliza una fuente heterogénea a una lista de diccionarios.

    Fuentes admitidas:
    - ``dict`` (un registro)
    - ``list``/``tuple``/iterable de ``dict``
    - ruta a ``.csv``, ``.tsv``, ``.txt``, ``.json``, ``.xlsx``/``.xlsm``
    - objetos tipo pandas ``DataFrame`` (con ``to_dict("records")``)
    - secuencia de secuencias (primera fila = cabeceras)

    Args:
        fuente: Origen de los datos.
        delimitador: Delimitador forzado para CSV/TXT (si es ``None``, se infiere).

    Returns:
        Lista de registros como diccionarios.

    Raises:
        ValueError: Si la fuente está vacía o no es interpretable.
        TypeError: Si el tipo de fuente no está soportado.
        FileNotFoundError: Si la ruta no existe.
    """
    if fuente is None:
        raise ValueError("La fuente de datos no puede ser None.")

    # pandas.DataFrame u objetos compatibles
    if hasattr(fuente, "to_dict") and callable(getattr(fuente, "to_dict")):
        try:
            registros = fuente.to_dict("records")
        except TypeError:
            registros = fuente.to_dict()
            if isinstance(registros, dict) and registros and isinstance(
                next(iter(registros.values())), (list, tuple)
            ):
                # formato columnar {col: [valores]}
                claves = list(registros.keys())
                n = len(registros[claves[0]]) if claves else 0
                return [
                    {k: registros[k][i] for k in claves}
                    for i in range(n)
                ]
        if not isinstance(registros, list):
            raise ValueError("No se pudo convertir el DataFrame a lista de registros.")
        return [_asegurar_dict(r, i) for i, r in enumerate(registros)]

    if isinstance(fuente, Mapping):
        return [dict(fuente)]

    if isinstance(fuente, (str, Path)):
        ruta = Path(fuente)
        if not ruta.exists():
            # Cadena JSON inline
            if isinstance(fuente, str) and fuente.strip().startswith(("[", "{")):
                return _desde_json_texto(fuente)
            raise FileNotFoundError(f"No existe el archivo: {ruta}")
        return _desde_archivo(ruta, delimitador=delimitador)

    if isinstance(fuente, Sequence) and not isinstance(fuente, (str, bytes, bytearray)):
        return _desde_secuencia(fuente)

    if isinstance(fuente, Iterable):
        return _desde_secuencia(list(fuente))

    raise TypeError(
        "Fuente no soportada. Use dict, list[dict], DataFrame, "
        "ruta (.csv/.txt/.json/.xlsx) o JSON en texto."
    )


def _asegurar_dict(registro: Any, indice: int) -> dict[str, Any]:
    if isinstance(registro, Mapping):
        return dict(registro)
    raise TypeError(f"El registro #{indice + 1} no es un diccionario.")


def _desde_secuencia(secuencia: Sequence[Any]) -> list[dict[str, Any]]:
    if len(secuencia) == 0:
        raise ValueError("La secuencia de datos está vacía.")

    primero = secuencia[0]
    if isinstance(primero, Mapping):
        return [_asegurar_dict(item, i) for i, item in enumerate(secuencia)]

    # Filas tipo tabla: [[headers], [row1], ...]
    if isinstance(primero, Sequence) and not isinstance(primero, (str, bytes)):
        cabeceras = [str(c).strip() for c in primero]
        if not cabeceras:
            raise ValueError("La tabla no tiene cabeceras.")
        registros: list[dict[str, Any]] = []
        for i, fila in enumerate(secuencia[1:], start=2):
            if not isinstance(fila, Sequence) or isinstance(fila, (str, bytes)):
                raise TypeError(f"La fila {i} no es una secuencia de valores.")
            valores = list(fila)
            if len(valores) < len(cabeceras):
                valores.extend([""] * (len(cabeceras) - len(valores)))
            registros.append(
                {cabeceras[j]: valores[j] for j in range(len(cabeceras))}
            )
        if not registros:
            raise ValueError("La tabla solo tiene cabeceras, sin filas de datos.")
        return registros

    raise TypeError(
        "Secuencia no reconocida. Espere list[dict] o tabla [cabeceras, fila1, ...]."
    )


def _desde_archivo(ruta: Path, *, delimitador: str | None) -> list[dict[str, Any]]:
    extension = ruta.suffix.lower()
    if extension in {".csv", ".tsv", ".txt"}:
        return _desde_delimitado(ruta, delimitador=delimitador, extension=extension)
    if extension == ".json":
        return _desde_json_texto(ruta.read_text(encoding="utf-8-sig"))
    if extension in {".xlsx", ".xlsm"}:
        return _desde_excel(ruta)
    raise ValueError(
        f"Extensión no soportada: {extension}. "
        "Use .csv, .tsv, .txt, .json, .xlsx o .xlsm."
    )


def _desde_delimitado(
    ruta: Path,
    *,
    delimitador: str | None,
    extension: str,
) -> list[dict[str, Any]]:
    texto = ruta.read_text(encoding="utf-8-sig")
    if not texto.strip():
        raise ValueError(f"El archivo está vacío: {ruta}")

    if delimitador is None:
        if extension == ".tsv":
            delimitador = "\t"
        else:
            try:
                dialecto = csv.Sniffer().sniff(texto[:2048], delimiters=";,\t|")
                delimitador = dialecto.delimiter
            except csv.Error:
                delimitador = ";" if texto.count(";") >= texto.count(",") else ","

    lector = csv.DictReader(texto.splitlines(), delimiter=delimitador)
    if not lector.fieldnames:
        raise ValueError("No se detectaron cabeceras en el archivo delimitado.")
    registros = [dict(fila) for fila in lector if any((v or "").strip() for v in fila.values())]
    if not registros:
        raise ValueError("El archivo delimitado no contiene filas de datos.")
    return registros


def _desde_json_texto(texto: str) -> list[dict[str, Any]]:
    try:
        datos = json.loads(texto)
    except json.JSONDecodeError as exc:
        raise ValueError(f"JSON inválido: {exc}") from exc
    return extraer_registros(datos)


def _desde_excel(ruta: Path) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook  # type: ignore[import-not-found]
    except ImportError as exc:
        raise ImportError(
            "Para importar Excel (.xlsx) instale openpyxl: pip install openpyxl"
        ) from exc

    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = libro.active
        filas = list(hoja.iter_rows(values_only=True))
    finally:
        libro.close()

    if not filas:
        raise ValueError(f"La hoja de Excel está vacía: {ruta}")

    # Normaliza None -> ""
    tabla = [
        ["" if celda is None else celda for celda in fila]
        for fila in filas
        if any(celda is not None and str(celda).strip() != "" for celda in fila)
    ]
    return _desde_secuencia(tabla)
